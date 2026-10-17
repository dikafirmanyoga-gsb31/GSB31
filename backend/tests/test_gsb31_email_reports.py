# GSB31 iteration 2: email reminder notifications (manual run, dedupe, disabled, cron webhook) + Excel report exports.
# NOTE: emails are REAL. This suite performs exactly 2 sending runs (manual + cron). Module-level functions, run with -n 0.
import os
import io
import time
import uuid
from datetime import timedelta, date

import pytest
import requests
from openpyxl import load_workbook
from dotenv import dotenv_values
from pymongo import MongoClient

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL").rstrip("/")
API = f"{BASE_URL}/api"
ENV = dotenv_values("/app/backend/.env")
CRON_SECRET = ENV["WEBHOOK_CRON_SECRET"]
DIR_ID, DIR_PW = "tengku.erens", "GSB31-Direktur!2026"
TEMP_PW = "GSB31-Sementara!26"
NEW_PW = "TestPass12345x"
RUN = uuid.uuid4().hex[:6]
STATE = {}
db = MongoClient(ENV["MONGO_URL"])[ENV["DB_NAME"]]


def login(ident, pw):
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"identifier": ident, "password": pw})
    return s, r


def staff_session(username):
    s, r = login(username, TEMP_PW)
    if r.status_code == 401:
        s, r = login(username, NEW_PW)
    assert r.status_code == 200, r.text
    if r.json()["must_change_password"]:
        c = s.post(f"{API}/auth/change-password", json={"current_password": TEMP_PW, "new_password": NEW_PW})
        assert c.status_code == 200, c.text
    return s


@pytest.fixture(scope="module")
def d():
    s, r = login(DIR_ID, DIR_PW)
    assert r.status_code == 200, r.text
    return s


@pytest.fixture(scope="module")
def putri():
    return staff_session("putri.anjani")


@pytest.fixture(scope="module")
def billy():
    return staff_session("billy")


def run_now(s):
    r = s.post(f"{API}/notifications/email-reminders/run")
    return r


def logs_count():
    return db.email_logs.count_documents({})


# ---------- RBAC on email endpoints ----------
def test_email_run_rbac(putri, billy):
    for s in (putri, billy):
        r = run_now(s)
        assert r.status_code == 403
        assert r.json()["detail"]["code"] == "DIRECTOR_ONLY"
        r = s.get(f"{API}/notifications/email-logs")
        assert r.status_code == 403
        assert r.json()["detail"]["code"] == "DIRECTOR_ONLY"


def test_email_logs_director(d):
    r = d.get(f"{API}/notifications/email-logs")
    assert r.status_code == 200
    j = r.json()
    assert isinstance(j["recipients"], list) and isinstance(j["logs"], list)
    assert any(x["email"] == "dikafirmanyoga@gmail.com" for x in j["recipients"])
    assert all("password_hash" not in x for x in j["recipients"])


# ---------- Cron auth ----------
def test_cron_unauthorized():
    assert requests.post(f"{API}/cron/reminder-emails").status_code == 401
    assert requests.post(f"{API}/cron/reminder-emails", headers={"Authorization": "Bearer wrong"}).status_code == 401
    assert requests.post(f"{API}/cron/reminder-emails", headers={"Authorization": CRON_SECRET}).status_code == 401


# ---------- Disabled ----------
def test_disabled_setting(d):
    r = d.put(f"{API}/settings", json={"email_reminders_enabled": False})
    assert r.status_code == 200 and r.json()["email_reminders_enabled"] is False
    try:
        r = run_now(d)
        assert r.status_code == 200
        assert r.json()["status"] == "disabled"
    finally:
        r = d.put(f"{API}/settings", json={"email_reminders_enabled": True})
        assert r.json()["email_reminders_enabled"] is True


# ---------- Setup: test recipient + doc due in 5 days ----------
def test_setup_recipient_and_doc(d):
    r = d.post(f"{API}/users", json={"name": "TEST_Email Rcpt", "username": f"test_mail_{RUN}", "email": "delivered@resend.dev",
                                     "role": "ops", "password": "TestPass12345x"})
    assert r.status_code == 200, r.text
    STATE["uid"] = r.json()["id"]
    # baseline: anything already due? (pre-existing TEST_ doc was emailed at H-14 by main agent)
    s = db.settings.find_one({"id": "app"}) or {}
    exp = (date.today() + timedelta(days=5)).isoformat()
    # use server tz date via documents days_left
    r = d.post(f"{API}/documents", json={"kind": "dokumen", "name": f"TEST_EmailDoc_{RUN}", "category": "Perizinan", "expiry_date": exp})
    assert r.status_code == 200, r.text
    STATE["doc"] = r.json()
    assert 3 <= r.json()["days_left"] <= 7


def test_manual_run_sends_digest(d):
    before = logs_count()
    r = run_now(d)
    assert r.status_code == 200, r.text
    j = r.json()
    print("RUN1", j)
    assert j["status"] == "done"
    assert j["items"] >= 1
    assert j["recipients"] == 2  # director + TEST recipient (only users with email)
    assert j["sent"] == 2 and j["failed"] == 0
    new = list(db.email_logs.find({}, {"_id": 0}).sort("at", -1).limit(logs_count() - before))
    assert len(new) == 2
    assert {x["to_email"] for x in new} == {"dikafirmanyoga@gmail.com", "delivered@resend.dev"}
    assert all(x["status"] == "sent" and x["trigger"] == "manual:tengku.erens" for x in new)
    rem = db.reminders.find_one({"source_id": STATE["doc"]["id"]})
    assert rem["emailed_levels"] == ["H-7"]
    # logs visible via API
    lg = d.get(f"{API}/notifications/email-logs").json()["logs"]
    assert lg[0]["status"] == "sent"


def test_second_run_nothing_due(d):
    before = logs_count()
    r = run_now(d)
    assert r.status_code == 200
    assert r.json()["status"] == "nothing_due"
    assert logs_count() == before


def test_expiry_change_resets_levels(d):
    doc = STATE["doc"]
    exp = (date.fromisoformat(doc["expiry_date"]) + timedelta(days=1)).isoformat()
    r = d.put(f"{API}/documents/{doc['id']}", json={"kind": "dokumen", "name": doc["name"], "category": "Perizinan", "expiry_date": exp})
    assert r.status_code == 200
    rem = db.reminders.find_one({"source_id": doc["id"]})
    assert rem["due_date"] == exp
    assert rem["emailed_levels"] == []


def test_cron_valid_and_idempotent():
    before = logs_count()
    wid = f"TEST_cron_{RUN}"
    h = {"Authorization": f"Bearer {CRON_SECRET}", "X-Webhook-Id": wid}
    t0 = time.time()
    r = requests.post(f"{API}/cron/reminder-emails", headers=h, json={})
    elapsed = time.time() - t0
    assert r.status_code == 200 and r.json()["accepted"] is True
    print("cron ack seconds", round(elapsed, 2))
    assert elapsed < 5
    for _ in range(20):
        time.sleep(1)
        if logs_count() - before >= 2:
            break
    assert logs_count() - before == 2
    new = list(db.email_logs.find({}, {"_id": 0}).sort("at", -1).limit(2))
    assert all(x["trigger"] == "cron" for x in new)
    assert db.cron_runs.find_one({"_id": wid}) is not None
    # same id again -> not processed
    mid = logs_count()
    r = requests.post(f"{API}/cron/reminder-emails", headers=h, json={})
    assert r.status_code == 200
    time.sleep(4)
    assert logs_count() == mid
    assert db.cron_runs.count_documents({"_id": wid}) == 1
    # new id but nothing due -> no emails
    h2 = {"Authorization": f"Bearer {CRON_SECRET}", "X-Webhook-Id": wid + "_b"}
    r = requests.post(f"{API}/cron/reminder-emails", headers=h2)
    assert r.status_code == 200
    time.sleep(4)
    assert logs_count() == mid


# ---------- Reports: finance ----------
def test_finance_report_setup(d):
    t = date.today()
    m1 = t.replace(day=1) - timedelta(days=15)
    txs = [("pemasukan", m1.isoformat(), "Penjualan Kopra", 1000000.5), ("pengeluaran", m1.isoformat(), "Transportasi", 250000),
           ("pengeluaran", t.isoformat(), "Operasional", 125000), ("pemasukan", t.isoformat(), "Pendapatan Lain", 999999)]
    ids = []
    for ty, dt, cat, amt in txs:
        r = d.post(f"{API}/finance/transactions", json={"type": ty, "date": dt, "category": cat, "amount": amt, "description": f"TEST_rep_{RUN}"})
        assert r.status_code == 200, r.text
        ids.append(r.json()["id"])
    r = d.post(f"{API}/finance/transactions/{ids[3]}/void", json={"reason": "TEST void"})
    assert r.status_code == 200
    STATE["from"], STATE["to"] = m1.replace(day=1).isoformat(), t.isoformat()


def read_xlsx(r):
    assert r.status_code == 200, r.text[:300]
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert "attachment" in r.headers["content-disposition"] and ".xlsx" in r.headers["content-disposition"]
    return load_workbook(io.BytesIO(r.content))


def table(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hi = next(i for i, r in enumerate(rows) if i >= 2 and r[0] and r[1] is not None and all(isinstance(x, str) or x is None for x in r) and i > 2)
    return rows[hi], [r for r in rows[hi + 1:] if any(v is not None for v in r)]


def test_finance_report_contents(d):
    p = {"date_from": STATE["from"], "date_to": STATE["to"]}
    wb = read_xlsx(d.get(f"{API}/reports/finance.xlsx", params=p))
    assert wb.sheetnames == ["Ringkasan", "Per Bulan", "Per Kategori", "Transaksi"]
    summ = d.get(f"{API}/finance/summary", params=p).json()
    assert summ["income"] == 1000000.5 and summ["expense"] == 375000 and summ["count"] == 3
    h, rows = table(wb["Ringkasan"])
    vals = {r[0]: r[1] for r in rows}
    assert vals["Total pemasukan"] == summ["income"]
    assert vals["Total pengeluaran"] == summ["expense"]
    assert vals["Selisih"] == summ["net"]
    assert vals["Jumlah transaksi"] == summ["count"]
    ws = wb["Ringkasan"]
    fmts = {ws.cell(row=i, column=1).value: ws.cell(row=i, column=2).number_format for i in range(1, 12)}
    print("RINGKASAN number formats", fmts)
    STATE["ringkasan_fmts"] = fmts
    h, rows = table(wb["Transaksi"])
    assert "Nominal" in h
    assert len(rows) == 3
    assert 999999 not in [r[h.index("Nominal")] for r in rows]  # voided excluded
    h, rows = table(wb["Per Bulan"])
    assert len(rows) == 2
    h, rows = table(wb["Per Kategori"])
    assert len(rows) == 3


def test_finance_report_number_formats():
    f = STATE["ringkasan_fmts"]
    assert "Rp" in f["Selisih"], f"Selisih should be currency, got {f['Selisih']}"
    assert "Rp" not in f["Jumlah transaksi"], f"Jumlah transaksi should be a plain count, got {f['Jumlah transaksi']}"


def test_finance_report_rbac(billy, putri):
    r = billy.get(f"{API}/reports/finance.xlsx")
    assert r.status_code == 403 and r.json()["detail"]["code"] == "FINANCE_RESTRICTED"
    r = putri.get(f"{API}/reports/finance.xlsx")
    assert r.status_code == 200
    assert requests.get(f"{API}/reports/finance.xlsx").status_code == 401


def test_report_invalid_date(d):
    r = d.get(f"{API}/reports/finance.xlsx", params={"date_from": "2026-13-45"})
    print("invalid date ->", r.status_code)
    assert r.status_code in (400, 422)


# ---------- Reports: kopra ----------
def test_kopra_report(d, billy, putri):
    t = date.today().isoformat()
    for st, q, price in (("selesai", 100, 5000), ("dibatalkan", 50, 5000)):
        r = d.post(f"{API}/kopra", json={"activity_type": "Pembelian", "date": t, "quantity": q, "unit": "kg", "status": st,
                                         "notes": f"TEST_kopra_{RUN}", "price_per_unit": price})
        assert r.status_code == 200, r.text
    p = {"date_from": t, "date_to": t}
    for s, fin in ((d, True), (putri, True), (billy, False)):
        wb = read_xlsx(s.get(f"{API}/reports/kopra.xlsx", params=p))
        all_vals = [v for ws in wb.worksheets for r in ws.iter_rows(values_only=True) for v in r if v is not None]
        for col in ("Harga/satuan", "Nilai", "Total nilai"):
            assert (col in all_vals) == fin, (col, fin)
        if fin:
            assert 500000 in all_vals  # total value of selesai row
        else:
            assert 500000 not in all_vals and 5000 not in all_vals
        ws = wb["Ringkasan"]
        rows = [r for r in ws.iter_rows(values_only=True)]
        pemb = [r for r in rows if r[0] == "Pembelian"]
        assert pemb and pemb[0][2] == 1 and pemb[0][3] == 100  # dibatalkan excluded from summary
        assert len([r for r in wb["Detail Kegiatan"].iter_rows(values_only=True) if r[0] == t]) == 2


def test_cleanup_test_recipient(d):
    r = d.patch(f"{API}/users/{STATE['uid']}", json={"is_active": False})
    assert r.status_code == 200
