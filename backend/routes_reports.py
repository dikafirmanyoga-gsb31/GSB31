from io import BytesIO
from datetime import date

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from core import db, require, has_perm, audit, today_jkt, get_settings, partner_map
from routes_finance import summarize, date_filter

router = APIRouter(prefix="/reports")
HEAD_FILL = PatternFill("solid", fgColor="0F172A")
HEAD_FONT = Font(bold=True, color="FFFFFF")
IDR = '"Rp" #,##0.00'


def sheet(ws, title_lines, headers, rows, money_cols=()):
    for i, line in enumerate(title_lines, 1):
        ws.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    hr = len(title_lines) + 2
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=hr, column=c, value=h)
        cell.fill, cell.font, cell.alignment = HEAD_FILL, HEAD_FONT, Alignment(vertical="center")
    for r, row in enumerate(rows, hr + 1):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            if c in money_cols:
                cell.number_format = IDR
    for c, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[c - 1])) for r in rows if r[c - 1] is not None] or [10])
        ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 10), 50)
    ws.freeze_panes = ws.cell(row=hr + 1, column=1)


def period_label(date_from, date_to):
    return f"Periode: {date_from or 'awal'} s/d {date_to or today_jkt().isoformat()}"


def xlsx_response(wb, name):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"', "Cache-Control": "no-store"})


def safe_date(s):
    try:
        return date.fromisoformat(s).isoformat() if s else ""
    except ValueError:
        raise HTTPException(400, "Format tanggal tidak valid (gunakan YYYY-MM-DD).")


@router.get("/finance.xlsx")
async def finance_report(date_from: str = "", date_to: str = "", user=Depends(require("finance"))):
    date_from, date_to = safe_date(date_from), safe_date(date_to)
    company = (await get_settings())["company"]["name"]
    s = await summarize(date_from, date_to)
    txs = [t async for t in db.transactions.find({"is_void": {"$ne": True}, **date_filter(date_from, date_to)}, {"_id": 0}).sort("date", 1)]
    names = await partner_map(t.get("partner_id") for t in txs)
    title = [f"Laporan Keuangan — {company}", period_label(date_from, date_to), f"Dibuat oleh {user['name']} pada {today_jkt().isoformat()}"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    sheet(ws, title, ["Keterangan", "Nilai"],
          [["Total pemasukan", s["income"]], ["Total pengeluaran", s["expense"]], ["Selisih", s["net"]], ["Jumlah transaksi", s["count"]]], (2,))
    ws.cell(row=len(title) + 6, column=2).number_format = "0"
    sheet(wb.create_sheet("Per Bulan"), title, ["Bulan", "Pemasukan", "Pengeluaran", "Selisih"],
          [[m["month"], m["pemasukan"], m["pengeluaran"], m["pemasukan"] - m["pengeluaran"]] for m in s["monthly"]], (2, 3, 4))
    sheet(wb.create_sheet("Per Kategori"), title, ["Jenis", "Kategori", "Nilai"],
          [[c["type"].capitalize(), c["category"], c["amount"]] for c in s["by_category"]], (3,))
    sheet(wb.create_sheet("Transaksi"), title, ["Tanggal", "Jenis", "Kategori", "Keterangan", "Mitra", "No. Referensi", "Nominal", "Dicatat oleh"],
          [[t["date"], t["type"].capitalize(), t["category"], t["description"], names.get(t.get("partner_id")), t.get("reference_no"),
            t["amount"], t.get("created_by_name")] for t in txs], (7,))
    await audit(user, "report_exported", "finance", None, detail=period_label(date_from, date_to))
    return xlsx_response(wb, f"GSB31_Laporan_Keuangan_{date_from or 'awal'}_{date_to or today_jkt().isoformat()}.xlsx")


@router.get("/kopra.xlsx")
async def kopra_report(date_from: str = "", date_to: str = "", user=Depends(require("ops"))):
    date_from, date_to = safe_date(date_from), safe_date(date_to)
    fin = has_perm(user, "finance")
    company = (await get_settings())["company"]["name"]
    items = [k async for k in db.kopra.find(date_filter(date_from, date_to), {"_id": 0}).sort("date", 1)]
    names = await partner_map(k.get("partner_id") for k in items)
    title = [f"Laporan Kegiatan Usaha Kopra — {company}", period_label(date_from, date_to), f"Dibuat oleh {user['name']} pada {today_jkt().isoformat()}"]
    agg = {}
    for k in items:
        if k["status"] == "dibatalkan":
            continue
        a = agg.setdefault((k["activity_type"], k["unit"]), {"n": 0, "q": 0.0, "v": 0.0})
        a["n"] += 1
        a["q"] += k["quantity"]
        a["v"] += k.get("total_value") or 0
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    head = ["Jenis kegiatan", "Satuan", "Jumlah kegiatan", "Total kuantitas"] + (["Total nilai"] if fin else [])
    sheet(ws, title + ["Tidak termasuk kegiatan berstatus Dibatalkan."], head,
          [[t, u, a["n"], round(a["q"], 2)] + ([round(a["v"], 2)] if fin else []) for (t, u), a in sorted(agg.items())], (5,) if fin else ())
    head = ["Tanggal", "Jenis", "Mitra", "Lokasi", "Kuantitas", "Satuan", "Status", "No. Referensi", "Catatan"] + (["Harga/satuan", "Nilai"] if fin else [])
    sheet(wb.create_sheet("Detail Kegiatan"), title, head,
          [[k["date"], k["activity_type"], names.get(k.get("partner_id")), k.get("location"), k["quantity"], k["unit"], k["status"].capitalize(),
            k.get("reference_no"), k.get("notes")] + ([k.get("price_per_unit"), k.get("total_value")] if fin else []) for k in items],
          (10, 11) if fin else ())
    await audit(user, "report_exported", "kopra", None, detail=period_label(date_from, date_to))
    return xlsx_response(wb, f"GSB31_Laporan_Kopra_{date_from or 'awal'}_{date_to or today_jkt().isoformat()}.xlsx")
